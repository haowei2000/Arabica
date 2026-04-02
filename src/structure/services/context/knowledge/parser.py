"""Document parser for various file types."""

import csv
import io
import logging
from typing import BinaryIO

from bs4 import BeautifulSoup
from docx import Document as DocxDocument
from pypdf import PdfReader

logger = logging.getLogger(__name__)


class DocumentParser:
    """Parse knowledge of various formats into plain text."""

    SUPPORTED_MIME_TYPES = {  # noqa: RUF012
        "text/plain": "txt",
        "text/markdown": "md",
        "application/pdf": "pdf",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
        "text/html": "html",
        "text/csv": "csv",
    }

    SUPPORTED_EXTENSIONS = {  # noqa: RUF012
        ".txt": "txt",
        ".md": "md",
        ".markdown": "md",
        ".pdf": "pdf",
        ".docx": "docx",
        ".html": "html",
        ".htm": "html",
        ".csv": "csv",
    }

    @staticmethod
    def _sanitize_text(text: str) -> str:
        """Remove null bytes and other invalid characters for PostgreSQL UTF-8.

        Args:
            text: Input text that may contain invalid characters.

        Returns:
            str: Sanitized text safe for PostgreSQL.
        """
        # Remove null bytes (0x00) which are invalid in PostgreSQL
        text = text.replace("\x00", "")
        # Remove other control characters except common whitespace
        sanitized = []
        for char in text:
            code = ord(char)
            # Keep: tab (9), newline (10), carriage return (13), and printable chars (32+)
            if code == 9 or code == 10 or code == 13 or code >= 32:
                sanitized.append(char)
        return "".join(sanitized)

    def parse(self, data: bytes | BinaryIO, mime_type: str) -> str:
        """Parse document content to plain text.

        Args:
            data: Document content as bytes or file-like object.
            mime_type: MIME type of the document.

        Returns:
            str: Extracted plain text content.

        Raises:
            ValueError: If MIME type is not supported.
        """
        file_type = self.SUPPORTED_MIME_TYPES.get(mime_type)
        if not file_type:
            raise ValueError(f"Unsupported MIME type: {mime_type}")

        # Convert bytes to stream if needed
        if isinstance(data, bytes):
            stream = io.BytesIO(data)
        else:
            stream = data
            stream.seek(0)

        parser_method = getattr(self, f"_parse_{file_type}", None)
        if not parser_method:
            raise ValueError(f"No parser available for type: {file_type}")

        text = parser_method(stream)
        # Sanitize text to remove invalid characters for PostgreSQL
        text = self._sanitize_text(text)
        logger.info(f"Parsed {file_type} document: {len(text)} characters")
        return text

    def parse_by_extension(self, data: bytes | BinaryIO, filename: str) -> str:
        """Parse document based on file extension.

        Args:
            data: Document content as bytes or file-like object.
            filename: Original filename with extension.

        Returns:
            str: Extracted plain text content.
        """
        ext = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
        file_type = self.SUPPORTED_EXTENSIONS.get(ext)

        if not file_type:
            raise ValueError(f"Unsupported file extension: {ext}")

        if isinstance(data, bytes):
            stream = io.BytesIO(data)
        else:
            stream = data
            stream.seek(0)

        parser_method = getattr(self, f"_parse_{file_type}", None)
        if not parser_method:
            raise ValueError(f"No parser available for type: {file_type}")

        text = parser_method(stream)
        # Sanitize text to remove invalid characters for PostgreSQL
        text = self._sanitize_text(text)
        logger.info(f"Parsed {file_type} document: {len(text)} characters")
        return text

    def _parse_txt(self, stream: BinaryIO) -> str:
        """Parse plain text file."""
        content = stream.read()
        # Try UTF-8 first, then fall back to other encodings
        for encoding in ["utf-8", "gbk", "gb2312", "latin-1"]:
            try:
                return content.decode(encoding)
            except UnicodeDecodeError:
                continue
        return content.decode("utf-8", errors="replace")

    def _parse_md(self, stream: BinaryIO) -> str:
        """Parse markdown file (return as plain text)."""
        return self._parse_txt(stream)

    def _parse_pdf(self, stream: BinaryIO) -> str:
        """Parse PDF file."""
        reader = PdfReader(stream)
        text_parts = []

        for page_num, page in enumerate(reader.pages):
            page_text = page.extract_text()
            if page_text:
                text_parts.append(page_text)
            else:
                logger.warning(f"No text extracted from PDF page {page_num + 1}")

        return "\n\n".join(text_parts)

    def _parse_docx(self, stream: BinaryIO) -> str:
        """Parse DOCX file."""
        doc = DocxDocument(stream)
        text_parts = []

        for para in doc.paragraphs:
            if para.text.strip():
                text_parts.append(para.text)

        # Also extract text from tables
        for table in doc.tables:
            for row in table.rows:
                row_text = []
                for cell in row.cells:
                    if cell.text.strip():
                        row_text.append(cell.text.strip())
                if row_text:
                    text_parts.append(" | ".join(row_text))

        return "\n\n".join(text_parts)

    def _parse_html(self, stream: BinaryIO) -> str:
        """Parse HTML file."""
        content = self._parse_txt(stream)
        soup = BeautifulSoup(content, "lxml")

        # Remove script and style elements
        for element in soup(["script", "style", "head", "meta", "link"]):
            element.decompose()

        # Get text content
        text = soup.get_text(separator="\n", strip=True)

        # Clean up multiple newlines
        lines = [line.strip() for line in text.split("\n") if line.strip()]
        return "\n".join(lines)

    def _parse_csv(self, stream: BinaryIO) -> str:
        """Parse CSV file."""
        content = self._parse_txt(stream)
        reader = csv.reader(io.StringIO(content))
        rows = []

        for row in reader:
            if any(cell.strip() for cell in row):
                rows.append(" | ".join(cell.strip() for cell in row))

        return "\n".join(rows)

    @classmethod
    def is_supported(cls, mime_type: str) -> bool:
        """Check if MIME type is supported.

        Args:
            mime_type: MIME type to check.

        Returns:
            bool: True if supported.
        """
        return mime_type in cls.SUPPORTED_MIME_TYPES

    @classmethod
    def get_supported_mime_types(cls) -> list[str]:
        """Get list of supported MIME types.

        Returns:
            list[str]: Supported MIME types.
        """
        return list(cls.SUPPORTED_MIME_TYPES.keys())
